import pandas as pd
import re
import unicodedata

from pathlib import Path
from difflib import SequenceMatcher

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# RUBAVU LOCATION CLEANING TOOL
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

DATA_DIR = BASE_DIR / "data"
BASELINE_DIR = DATA_DIR / "baseline"
PROCESSED_DIR = DATA_DIR / "processed"

MASTER_FILE = BASELINE_DIR / "RubavunewAugust.xlsx"

if not MASTER_FILE.exists():
    MASTER_FILE = BASE_DIR / "RubavunewAugust.xlsx"

OUTPUT_FILE = BASE_DIR / "location_mapping.xlsx"


# ============================================================
# OFFICIAL RUBAVU STRUCTURE
# ============================================================

EXPECTED_DISTRICTS = 1
EXPECTED_SECTORS = 12
EXPECTED_CELLS = 80
EXPECTED_VILLAGES = 519


# ============================================================
# LOCATION COLUMNS
# ============================================================

LOCATION_COLUMNS = [
    "district",
    "sector",
    "cell",
    "village"
]


# ============================================================
# COLUMN NORMALIZATION
# ============================================================

def clean_column_name(value):

    value = str(value).strip().lower()

    value = re.sub(
        r"[\s\-/]+",
        "_",
        value
    )

    value = re.sub(
        r"[^a-z0-9_]",
        "",
        value
    )

    return value


def normalize_columns(df):

    df = df.copy()

    df.columns = [
        clean_column_name(c)
        for c in df.columns
    ]

    return df


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(value):

    if pd.isna(value):
        return ""

    value = str(value).strip()

    if value.lower() in [
        "nan",
        "none",
        "null"
    ]:
        return ""

    # Unicode normalization
    value = unicodedata.normalize(
        "NFKD",
        value
    )

    # Remove accents
    value = "".join(
        c
        for c in value
        if not unicodedata.combining(c)
    )

    # Lowercase
    value = value.casefold()

    # Normalize punctuation
    value = re.sub(
        r"[^a-z0-9\s]",
        " ",
        value
    )

    # Multiple spaces
    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


# ============================================================
# COMPARISON KEY
# ============================================================

def comparison_key(value):

    value = normalize_text(value)

    # Remove spaces for comparison
    return value.replace(
        " ",
        ""
    )


# ============================================================
# SIMILARITY
# ============================================================

def similarity(a, b):

    a = comparison_key(a)
    b = comparison_key(b)

    if not a or not b:
        return 0

    return SequenceMatcher(
        None,
        a,
        b
    ).ratio()


# ============================================================
# READ EXCEL
# ============================================================

def read_excel(file):

    try:

        df = pd.read_excel(
            file
        )

        return normalize_columns(
            df
        )

    except Exception as e:

        print(
            f"ERROR reading {file.name}: {e}"
        )

        return pd.DataFrame()


# ============================================================
# LOAD OFFICIAL MASTER
# ============================================================

def load_master():

    print()
    print("=" * 70)
    print("LOADING OFFICIAL RUBAVU MASTER")
    print("=" * 70)

    print()
    print(
        f"Master file: {MASTER_FILE}"
    )

    if not MASTER_FILE.exists():

        print()
        print("ERROR: Master file not found.")

        print(
            f"Expected: {MASTER_FILE}"
        )

        return pd.DataFrame()

    master = read_excel(
        MASTER_FILE
    )

    if master.empty:

        print(
            "ERROR: Master file is empty."
        )

        return pd.DataFrame()

    # --------------------------------------------------------
    # Required columns
    # --------------------------------------------------------

    required = [
        "sector",
        "cell",
        "village"
    ]

    missing = [
        c
        for c in required
        if c not in master.columns
    ]

    if missing:

        print()
        print(
            "ERROR: Master is missing:"
        )

        print(
            missing
        )

        print()
        print(
            "Columns found:"
        )

        print(
            list(master.columns)
        )

        return pd.DataFrame()

    # --------------------------------------------------------
    # District is optional
    # --------------------------------------------------------

    if "district" not in master.columns:

        master["district"] = "Rubavu"

    # --------------------------------------------------------
    # Clean values
    # --------------------------------------------------------

    for col in LOCATION_COLUMNS:

        master[col] = (
            master[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

    # --------------------------------------------------------
    # Remove empty geographic rows
    # --------------------------------------------------------

    master = master[
        (master["sector"] != "")
        &
        (master["cell"] != "")
        &
        (master["village"] != "")
    ].copy()

    # --------------------------------------------------------
    # Create comparison keys
    # --------------------------------------------------------

    for col in LOCATION_COLUMNS:

        master[
            f"{col}_key"
        ] = master[col].apply(
            comparison_key
        )

    # --------------------------------------------------------
    # Remove duplicate geographic records
    # --------------------------------------------------------

    master = master.drop_duplicates(
        subset=[
            "district_key",
            "sector_key",
            "cell_key",
            "village_key"
        ]
    )

    # ========================================================
    # MASTER COUNTS
    # ========================================================

    districts = (
        master[
            "district_key"
        ].nunique()
    )

    sectors = (
        master[
            "sector_key"
        ].nunique()
    )

    cells = (
        master[
            [
                "sector_key",
                "cell_key"
            ]
        ]
        .drop_duplicates()
        .shape[0]
    )

    villages = (
        master[
            [
                "sector_key",
                "cell_key",
                "village_key"
            ]
        ]
        .drop_duplicates()
        .shape[0]
    )

    print()
    print("MASTER STRUCTURE")
    print()

    print(
        f"Districts : {districts}"
    )

    print(
        f"Sectors   : {sectors}"
    )

    print(
        f"Cells     : {cells}"
    )

    print(
        f"Villages  : {villages}"
    )

    # --------------------------------------------------------
    # Validation
    # --------------------------------------------------------

    if districts == EXPECTED_DISTRICTS:
        print(
            "✓ 1 district confirmed"
        )
    else:
        print(
            f"⚠ Expected {EXPECTED_DISTRICTS} "
            f"district but found {districts}"
        )

    if sectors == EXPECTED_SECTORS:
        print(
            "✓ 12 sectors confirmed"
        )
    else:
        print(
            f"⚠ Expected {EXPECTED_SECTORS} "
            f"sectors but found {sectors}"
        )

    if cells == EXPECTED_CELLS:
        print(
            "✓ 80 cells confirmed"
        )
    else:
        print(
            f"⚠ Expected {EXPECTED_CELLS} "
            f"cells but found {cells}"
        )

    if villages == EXPECTED_VILLAGES:
        print(
            "✓ 519 villages confirmed"
        )
    else:
        print(
            f"⚠ Expected {EXPECTED_VILLAGES} "
            f"villages but found {villages}"
        )

    return master


# ============================================================
# FIND SOURCE FILES
# ============================================================

def find_source_files():

    files = []

    # --------------------------------------------------------
    # Processed files
    # --------------------------------------------------------

    if PROCESSED_DIR.exists():

        files.extend(
            PROCESSED_DIR.glob(
                "*_matched.xlsx"
            )
        )

        files.extend(
            PROCESSED_DIR.glob(
                "*_processed.xlsx"
            )
        )

        files.extend(
            PROCESSED_DIR.glob(
                "*.xls"
            )
        )

    # --------------------------------------------------------
    # Remove master file if accidentally included
    # --------------------------------------------------------

    files = [
        f
        for f in files
        if f.resolve()
        != MASTER_FILE.resolve()
    ]

    # --------------------------------------------------------
    # Remove duplicates
    # --------------------------------------------------------

    files = list(
        dict.fromkeys(files)
    )

    return sorted(
        files
    )


# ============================================================
# PREPARE SOURCE
# ============================================================

def prepare_source(df):

    df = df.copy()

    # --------------------------------------------------------
    # District is optional
    # --------------------------------------------------------

    if "district" not in df.columns:

        df["district"] = ""

    # --------------------------------------------------------
    # Sector / Cell / Village required
    # --------------------------------------------------------

    for col in [
        "sector",
        "cell",
        "village"
    ]:

        if col not in df.columns:

            df[col] = ""

    # --------------------------------------------------------
    # Clean
    # --------------------------------------------------------

    for col in LOCATION_COLUMNS:

        df[col] = (
            df[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

        df[
            f"{col}_key"
        ] = df[col].apply(
            comparison_key
        )

    return df


# ============================================================
# EXACT MATCH
# ============================================================

def exact_match(
    source_value,
    candidates
):

    source_key = comparison_key(
        source_value
    )

    if not source_key:

        return None

    for candidate in candidates:

        if comparison_key(
            candidate
        ) == source_key:

            return candidate

    return None


# ============================================================
# FUZZY MATCH
# ============================================================

def fuzzy_match(
    source_value,
    candidates,
    threshold=0.90
):

    if not source_value:

        return (
            None,
            0
        )

    best_candidate = None
    best_score = 0

    for candidate in candidates:

        score = similarity(
            source_value,
            candidate
        )

        if score > best_score:

            best_score = score
            best_candidate = candidate

    if (
        best_candidate
        and
        best_score >= threshold
    ):

        return (
            best_candidate,
            best_score
        )

    return (
        None,
        best_score
    )


# ============================================================
# MATCH DISTRICT
# ============================================================

def match_district(
    source_district,
    master
):

    candidates = (
        master[
            "district"
        ]
        .drop_duplicates()
        .tolist()
    )

    # If source has no district,
    # use official master district.

    if not source_district:

        if len(candidates) == 1:

            return (
                candidates[0],
                "MASTER_DEFAULT",
                100
            )

        return (
            "",
            "MISSING",
            0
        )

    # Exact
    exact = exact_match(
        source_district,
        candidates
    )

    if exact:

        return (
            exact,
            "EXACT",
            100
        )

    # Fuzzy
    match, score = fuzzy_match(
        source_district,
        candidates,
        threshold=0.95
    )

    if match:

        return (
            match,
            "FUZZY_REVIEW",
            round(
                score * 100,
                2
            )
        )

    return (
        "",
        "UNMATCHED",
        round(
            score * 100,
            2
        )
    )


# ============================================================
# MATCH SECTOR
# ============================================================

def match_sector(
    source_sector,
    master
):

    candidates = (
        master[
            "sector"
        ]
        .drop_duplicates()
        .tolist()
    )

    # Exact normalized
    exact = exact_match(
        source_sector,
        candidates
    )

    if exact:

        return (
            exact,
            "EXACT",
            100
        )

    # Fuzzy
    match, score = fuzzy_match(
        source_sector,
        candidates,
        threshold=0.90
    )

    if match:

        return (
            match,
            "FUZZY_REVIEW",
            round(
                score * 100,
                2
            )
        )

    return (
        "",
        "UNMATCHED",
        round(
            score * 100,
            2
        )
    )


# ============================================================
# MATCH CELL
# ============================================================

def match_cell(
    source_sector,
    source_cell,
    master
):

    # --------------------------------------------------------
    # Find sector
    # --------------------------------------------------------

    sector_match, sector_status, sector_score = (
        match_sector(
            source_sector,
            master
        )
    )

    if not sector_match:

        return (
            "",
            "UNMATCHED_PARENT_SECTOR",
            0
        )

    sector_key = comparison_key(
        sector_match
    )

    # --------------------------------------------------------
    # ONLY cells inside this sector
    # --------------------------------------------------------

    candidates = (
        master[
            master["sector_key"]
            ==
            sector_key
        ]["cell"]
        .drop_duplicates()
        .tolist()
    )

    # --------------------------------------------------------
    # Exact
    # --------------------------------------------------------

    exact = exact_match(
        source_cell,
        candidates
    )

    if exact:

        return (
            exact,
            "EXACT",
            100
        )

    # --------------------------------------------------------
    # Fuzzy
    # --------------------------------------------------------

    match, score = fuzzy_match(
        source_cell,
        candidates,
        threshold=0.90
    )

    if match:

        return (
            match,
            "FUZZY_REVIEW",
            round(
                score * 100,
                2
            )
        )

    return (
        "",
        "UNMATCHED",
        round(
            score * 100,
            2
        )
    )


# ============================================================
# MATCH VILLAGE
# ============================================================

def match_village(
    source_sector,
    source_cell,
    source_village,
    master
):

    # --------------------------------------------------------
    # Sector
    # --------------------------------------------------------

    sector_match, _, _ = (
        match_sector(
            source_sector,
            master
        )
    )

    if not sector_match:

        return (
            "",
            "UNMATCHED_PARENT_SECTOR",
            0
        )

    sector_key = comparison_key(
        sector_match
    )

    # --------------------------------------------------------
    # Cell
    # --------------------------------------------------------

    cell_candidates = (
        master[
            master["sector_key"]
            ==
            sector_key
        ]["cell"]
        .drop_duplicates()
        .tolist()
    )

    cell_match = exact_match(
        source_cell,
        cell_candidates
    )

    # If exact cell isn't found,
    # fuzzy match inside the sector.

    if not cell_match:

        cell_match, cell_score = fuzzy_match(
            source_cell,
            cell_candidates,
            threshold=0.90
        )

    if not cell_match:

        return (
            "",
            "UNMATCHED_PARENT_CELL",
            0
        )

    cell_key = comparison_key(
        cell_match
    )

    # --------------------------------------------------------
    # ONLY villages inside this cell
    # --------------------------------------------------------

    candidates = (
        master[
            (
                master["sector_key"]
                ==
                sector_key
            )
            &
            (
                master["cell_key"]
                ==
                cell_key
            )
        ]["village"]
        .drop_duplicates()
        .tolist()
    )

    # --------------------------------------------------------
    # Exact normalized match
    # --------------------------------------------------------

    exact = exact_match(
        source_village,
        candidates
    )

    if exact:

        return (
            exact,
            "EXACT",
            100
        )

    # --------------------------------------------------------
    # Fuzzy match
    # --------------------------------------------------------

    match, score = fuzzy_match(
        source_village,
        candidates,
        threshold=0.90
    )

    if match:

        return (
            match,
            "FUZZY_REVIEW",
            round(
                score * 100,
                2
            )
        )

    return (
        "",
        "UNMATCHED",
        round(
            score * 100,
            2
        )
    )


# ============================================================
# MATCH ONE SOURCE FILE
# ============================================================

def build_source_mapping(
    source,
    master,
    source_file
):

    if source.empty:

        return pd.DataFrame()

    # --------------------------------------------------------
    # Prepare source
    # --------------------------------------------------------

    source = prepare_source(
        source
    )

    # --------------------------------------------------------
    # Geographic records
    #
    # IMPORTANT:
    # district is NOT required
    # --------------------------------------------------------

    locations = (
        source[
            [
                "district",
                "sector",
                "cell",
                "village"
            ]
        ]
        .drop_duplicates()
    )

    print()
    print(
        f"Matching: {source_file}"
    )

    print(
        f"Unique geographic records: "
        f"{len(locations):,}"
    )

    rows = []

    # --------------------------------------------------------
    # Loop locations
    # --------------------------------------------------------

    for _, row in locations.iterrows():

        original_district = row[
            "district"
        ]

        original_sector = row[
            "sector"
        ]

        original_cell = row[
            "cell"
        ]

        original_village = row[
            "village"
        ]

        # ----------------------------------------------------
        # District
        # ----------------------------------------------------

        (
            district_master,
            district_status,
            district_score
        ) = match_district(
            original_district,
            master
        )

        # ----------------------------------------------------
        # Sector
        # ----------------------------------------------------

        (
            sector_master,
            sector_status,
            sector_score
        ) = match_sector(
            original_sector,
            master
        )

        # ----------------------------------------------------
        # Cell
        # ----------------------------------------------------

        (
            cell_master,
            cell_status,
            cell_score
        ) = match_cell(
            original_sector,
            original_cell,
            master
        )

        # ----------------------------------------------------
        # Village
        # ----------------------------------------------------

        (
            village_master,
            village_status,
            village_score
        ) = match_village(
            original_sector,
            original_cell,
            original_village,
            master
        )

        # ----------------------------------------------------
        # Overall status
        # ----------------------------------------------------

        if (
            sector_status == "EXACT"
            and
            cell_status == "EXACT"
            and
            village_status == "EXACT"
        ):

            overall = "AUTO_APPROVED"

        elif (
            village_status == "FUZZY_REVIEW"
            or
            cell_status == "FUZZY_REVIEW"
            or
            sector_status == "FUZZY_REVIEW"
        ):

            overall = "REVIEW"

        elif (
            village_status.startswith(
                "UNMATCHED"
            )
            or
            cell_status.startswith(
                "UNMATCHED"
            )
            or
            sector_status == "UNMATCHED"
        ):

            overall = "UNMATCHED"

        else:

            overall = "REVIEW"

        # ----------------------------------------------------
        # Record
        # ----------------------------------------------------

        rows.append({

            "Source_File":
                source_file,

            "District_Original":
                original_district,

            "Sector_Original":
                original_sector,

            "Cell_Original":
                original_cell,

            "Village_Original":
                original_village,

            "District_Master":
                district_master,

            "Sector_Master":
                sector_master,

            "Cell_Master":
                cell_master,

            "Village_Master":
                village_master,

            "District_Status":
                district_status,

            "Sector_Status":
                sector_status,

            "Cell_Status":
                cell_status,

            "Village_Status":
                village_status,

            "District_Score":
                district_score,

            "Sector_Score":
                sector_score,

            "Cell_Score":
                cell_score,

            "Village_Score":
                village_score,

            "Overall_Status":
                overall,

            "Approved":
                (
                    "YES"
                    if overall
                    == "AUTO_APPROVED"
                    else ""
                )
        })

    return pd.DataFrame(
        rows
    )


# ============================================================
# MASTER SHEET
# ============================================================

def create_master_sheet(master):

    result = master[
        [
            "district",
            "sector",
            "cell",
            "village"
        ]
    ].drop_duplicates()

    result = result.sort_values(
        [
            "sector",
            "cell",
            "village"
        ]
    )

    result.columns = [
        "District",
        "Sector",
        "Cell",
        "Village"
    ]

    result.reset_index(
        drop=True,
        inplace=True
    )

    return result


# ============================================================
# APPROVED MAPPING
# ============================================================

def create_approved_mapping(
    mappings
):

    if mappings.empty:

        return pd.DataFrame()

    result = mappings[
        mappings[
            "Overall_Status"
        ]
        ==
        "AUTO_APPROVED"
    ].copy()

    result = result[
        [
            "Sector_Original",
            "Cell_Original",
            "Village_Original",

            "Sector_Master",
            "Cell_Master",
            "Village_Master"
        ]
    ]

    result.columns = [
        "Original_Sector",
        "Original_Cell",
        "Original_Village",

        "Canonical_Sector",
        "Canonical_Cell",
        "Canonical_Village"
    ]

    return result.drop_duplicates()


# ============================================================
# REVIEW SHEET
# ============================================================

def create_review_sheet(
    mappings
):

    if mappings.empty:

        return pd.DataFrame()

    result = mappings[
        mappings[
            "Overall_Status"
        ]
        .isin(
            [
                "REVIEW",
                "UNMATCHED"
            ]
        )
    ].copy()

    return result


# ============================================================
# SPELLING DUPLICATES
# ============================================================

def detect_spelling_duplicates(
    mappings
):

    if mappings.empty:

        return pd.DataFrame()

    rows = []

    # --------------------------------------------------------
    # Only compare records that have a canonical match
    # --------------------------------------------------------

    for level in [
        "Sector",
        "Cell",
        "Village"
    ]:

        original_col = (
            f"{level}_Original"
        )

        master_col = (
            f"{level}_Master"
        )

        if (
            original_col
            not in mappings.columns
        ):
            continue

        if (
            master_col
            not in mappings.columns
        ):
            continue

        subset = mappings[
            mappings[
                master_col
            ].astype(str).str.strip()
            != ""
        ][
            [
                original_col,
                master_col
            ]
        ].drop_duplicates()

        for _, row in subset.iterrows():

            original = str(
                row[
                    original_col
                ]
            ).strip()

            canonical = str(
                row[
                    master_col
                ]
            ).strip()

            if not original:

                continue

            if not canonical:

                continue

            # Different written forms
            if normalize_text(
                original
            ) != normalize_text(
                canonical
            ):

                score = similarity(
                    original,
                    canonical
                )

                rows.append({

                    "Level":
                        level,

                    "Original_Name":
                        original,

                    "Canonical_Name":
                        canonical,

                    "Similarity":
                        round(
                            score * 100,
                            2
                        ),

                    "Action":
                        (
                            "AUTO_NORMALIZED"
                            if score >= 0.90
                            else "REVIEW"
                        )
                })

    if not rows:

        return pd.DataFrame(
            columns=[
                "Level",
                "Original_Name",
                "Canonical_Name",
                "Similarity",
                "Action"
            ]
        )

    result = pd.DataFrame(
        rows
    )

    return result.sort_values(
        [
            "Level",
            "Similarity"
        ],
        ascending=[
            True,
            False
        ]
    )


# ============================================================
# SUMMARY
# ============================================================

def create_summary(
    master,
    mappings,
    duplicates
):

    master_sectors = (
        master[
            "sector_key"
        ].nunique()
    )

    master_cells = (
        master[
            [
                "sector_key",
                "cell_key"
            ]
        ]
        .drop_duplicates()
        .shape[0]
    )

    master_villages = (
        master[
            [
                "sector_key",
                "cell_key",
                "village_key"
            ]
        ]
        .drop_duplicates()
        .shape[0]
    )

    if mappings.empty:

        auto = 0
        review = 0
        unmatched = 0

    else:

        auto = (
            mappings[
                "Overall_Status"
            ]
            ==
            "AUTO_APPROVED"
        ).sum()

        review = (
            mappings[
                "Overall_Status"
            ]
            ==
            "REVIEW"
        ).sum()

        unmatched = (
            mappings[
                "Overall_Status"
            ]
            ==
            "UNMATCHED"
        ).sum()

    return pd.DataFrame({

        "Metric": [

            "Expected Districts",
            "Master Districts",

            "Expected Sectors",
            "Master Sectors",

            "Expected Cells",
            "Master Cells",

            "Expected Villages",
            "Master Villages",

            "Source Geographic Records",

            "Auto Approved",

            "Needs Review",

            "Unmatched",

            "Spelling Variations Detected"

        ],

        "Value": [

            EXPECTED_DISTRICTS,
            master[
                "district_key"
            ].nunique(),

            EXPECTED_SECTORS,
            master_sectors,

            EXPECTED_CELLS,
            master_cells,

            EXPECTED_VILLAGES,
            master_villages,

            len(mappings),

            auto,

            review,

            unmatched,

            len(duplicates)

        ]
    })


# ============================================================
# WRITE EXCEL
# ============================================================

def write_excel(
    master,
    mappings,
    duplicates
):

    master_sheet = create_master_sheet(
        master
    )

    approved = create_approved_mapping(
        mappings
    )

    review = create_review_sheet(
        mappings
    )

    summary = create_summary(
        master,
        mappings,
        duplicates
    )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        master_sheet.to_excel(
            writer,
            sheet_name="MASTER_519",
            index=False
        )

        mappings.to_excel(
            writer,
            sheet_name="ALL_MAPPING",
            index=False
        )

        approved.to_excel(
            writer,
            sheet_name="APPROVED_MAPPING",
            index=False
        )

        review.to_excel(
            writer,
            sheet_name="NEEDS_REVIEW",
            index=False
        )

        duplicates.to_excel(
            writer,
            sheet_name="SPELLING_DUPLICATES",
            index=False
        )

        summary.to_excel(
            writer,
            sheet_name="SUMMARY",
            index=False
        )


# ============================================================
# FORMAT EXCEL
# ============================================================

def format_excel():

    wb = load_workbook(
        OUTPUT_FILE
    )

    # --------------------------------------------------------
    # Styles
    # --------------------------------------------------------

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    review_fill = PatternFill(
        "solid",
        fgColor="FFF2CC"
    )

    approved_fill = PatternFill(
        "solid",
        fgColor="E2F0D9"
    )

    unmatched_fill = PatternFill(
        "solid",
        fgColor="F4CCCC"
    )

    duplicate_fill = PatternFill(
        "solid",
        fgColor="DDEBF7"
    )

    # --------------------------------------------------------
    # Sheets
    # --------------------------------------------------------

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        # Header
        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Column widths
        for column_cells in ws.columns:

            max_length = 0

            col_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                try:

                    length = len(
                        str(
                            cell.value
                        )
                    )

                    if length > max_length:
                        max_length = length

                except Exception:
                    pass

            ws.column_dimensions[
                col_letter
            ].width = min(
                max_length + 2,
                50
            )

    # --------------------------------------------------------
    # Status coloring
    # --------------------------------------------------------

    for ws in wb.worksheets:

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                value = str(
                    cell.value
                ).upper()

                if value == "REVIEW":

                    cell.fill = review_fill

                elif value in [
                    "YES",
                    "AUTO_APPROVED"
                ]:

                    cell.fill = approved_fill

                elif value == "UNMATCHED":

                    cell.fill = unmatched_fill

                elif value == "AUTO_NORMALIZED":

                    cell.fill = duplicate_fill

    wb.save(
        OUTPUT_FILE
    )


# ============================================================
# PRINT REPORT
# ============================================================

def print_report(
    master,
    mappings,
    duplicates
):

    print()
    print("=" * 70)
    print("LOCATION CLEANING REPORT")
    print("=" * 70)

    print()

    print("OFFICIAL MASTER")
    print()

    print(
        f"Districts : "
        f"{master['district_key'].nunique()}"
    )

    print(
        f"Sectors   : "
        f"{master['sector_key'].nunique()}"
    )

    print(
        f"Cells     : "
        f"{master[['sector_key', 'cell_key']].drop_duplicates().shape[0]}"
    )

    print(
        f"Villages  : "
        f"{master[['sector_key', 'cell_key', 'village_key']].drop_duplicates().shape[0]}"
    )

    print()

    print("MAPPING")
    print()

    if not mappings.empty:

        auto = (
            mappings[
                "Overall_Status"
            ]
            ==
            "AUTO_APPROVED"
        ).sum()

        review = (
            mappings[
                "Overall_Status"
            ]
            ==
            "REVIEW"
        ).sum()

        unmatched = (
            mappings[
                "Overall_Status"
            ]
            ==
            "UNMATCHED"
        ).sum()

    else:

        auto = 0
        review = 0
        unmatched = 0

    print(
        f"Source locations : "
        f"{len(mappings):,}"
    )

    print(
        f"Auto approved    : "
        f"{auto:,}"
    )

    print(
        f"Needs review     : "
        f"{review:,}"
    )

    print(
        f"Unmatched        : "
        f"{unmatched:,}"
    )

    print()

    print(
        f"Spelling variations detected : "
        f"{len(duplicates):,}"
    )

    print()

    # --------------------------------------------------------
    # Show important examples
    # --------------------------------------------------------

    if not duplicates.empty:

        print(
            "TOP SPELLING VARIATIONS"
        )

        print()

        display_columns = [
            "Level",
            "Original_Name",
            "Canonical_Name",
            "Similarity",
            "Action"
        ]

        print(
            duplicates[
                display_columns
            ]
            .head(30)
            .to_string(
                index=False
            )
        )

    print()
    print("=" * 70)

    print(
        "Excel created:"
    )

    print(
        OUTPUT_FILE
    )

    print("=" * 70)


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 70)
    print("RUBAVU LOCATION MASTER CLEANER")
    print("=" * 70)

    print()
    print(
        "Official master:"
    )

    print(
        MASTER_FILE
    )

    print()
    print(
        "Accepted official structure:"
    )

    print(
        "1 District | 12 Sectors | 80 Cells | 519 Villages"
    )

    # ========================================================
    # MASTER
    # ========================================================

    master = load_master()

    if master.empty:

        return

    # ========================================================
    # SOURCE FILES
    # ========================================================

    source_files = find_source_files()

    print()

    print(
        f"Progress files found: "
        f"{len(source_files)}"
    )

    if not source_files:

        print()

        print(
            "No processed files found."
        )

        print(
            f"Directory: {PROCESSED_DIR}"
        )

        return

    # ========================================================
    # MATCH ALL FILES
    # ========================================================

    all_mappings = []

    for file in source_files:

        print()
        print("-" * 70)

        print(
            f"READING: {file.name}"
        )

        source = read_excel(
            file
        )

        if source.empty:

            print(
                "Skipped: empty/unreadable."
            )

            continue

        mapping = build_source_mapping(
            source,
            master,
            file.name
        )

        if not mapping.empty:

            all_mappings.append(
                mapping
            )

    # ========================================================
    # COMBINE
    # ========================================================

    if not all_mappings:

        print()
        print(
            "ERROR: No mapping records created."
        )

        return

    mappings = pd.concat(
        all_mappings,
        ignore_index=True
    )

    # ========================================================
    # DETECT SPELLING DUPLICATES
    # ========================================================

    print()
    print("=" * 70)

    print(
        "DETECTING SPELLING DUPLICATES"
    )

    print("=" * 70)

    duplicates = detect_spelling_duplicates(
        mappings
    )

    print()
    print(
        f"Spelling variations found: "
        f"{len(duplicates):,}"
    )

    # ========================================================
    # CREATE EXCEL
    # ========================================================

    print()
    print("=" * 70)

    print(
        "CREATING EXCEL MAPPING"
    )

    print("=" * 70)

    write_excel(
        master,
        mappings,
        duplicates
    )

    print(
        "Formatting Excel workbook..."
    )

    format_excel()

    # ========================================================
    # REPORT
    # ========================================================

    print_report(
        master,
        mappings,
        duplicates
    )

    print()
    print(
        "DONE."
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()